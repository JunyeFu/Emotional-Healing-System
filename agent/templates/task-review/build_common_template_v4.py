"""Build the reusable Chinese task-package form and its Word template."""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / 'human/templates'
NAME = '通用中文逐包任务模板_v4.0'


def table(doc, headers, rows, widths):
    grid = doc.add_table(rows=1, cols=len(headers))
    grid.autofit = False
    for column, width in zip(grid.columns, widths):
        column.width = Inches(width)
    for cell, value, width in zip(grid.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.text = value
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'E9EEF1')
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement('w:tblHeader')
    grid.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells = grid.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.width = Inches(width)
            cell.text = value
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement(f'w:{side}')
        for key, value in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
            border.set(qn(f'w:{key}'), value)
        borders.append(border)
    grid._tbl.tblPr.append(borders)
    for row in grid.rows:
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        for cell in row.cells:
            cell.vertical_alignment = 1
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                item = OxmlElement(f'w:{side}')
                item.set(qn('w:w'), '90')
                item.set(qn('w:type'), 'dxa')
                margins.append(item)
            cell._tc.get_or_add_tcPr().append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    run.font.size = Pt(10)


def heading(doc, title, text):
    doc.add_heading(title, 1)
    doc.add_paragraph(text)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.65)
    sec.left_margin = sec.right_margin = Inches(.75)
    for name, size in [('Normal', 11), ('Title', 21), ('Heading 1', 13), ('Heading 2', 11)]:
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.08
    for name in ('Heading 1', 'Heading 2'):
        doc.styles[name].paragraph_format.space_before = Pt(9)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    header = sec.header.paragraphs[0]
    header.text = '逐包任务工作单   通用版 4.0'
    header.runs[0].font.size = Pt(9)
    footer = sec.footer.paragraphs[0]
    footer.text = '任务编号 [待填]   文档版本 [待填]                                      第 '
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    footer.add_run(' 页')
    for run in footer.runs:
        run.font.size = Pt(9)

    doc.add_paragraph('通用中文逐包任务模板', 'Title')
    doc.add_paragraph('适用于科研、编程开发、数学建模、课程作业与投资研究。一包一份，领取时定义，执行时增量填写，复核后交接；正文只记录本包的必要事实与决定。')
    doc.add_paragraph('填写约定  方括号为待填项，未执行或不适用分别注明。文档用于人读，代码、数据与命令以链接为准。小任务可压缩；独立复核、签收和发布仅按任务要求填写。')
    table(doc, ['基本信息', '填写内容'], [
        ['项目与任务', '[项目名] / [任务编号]【领域】[动词＋交付对象]'],
        ['当前状态与结论', '[沿用项目状态体系]；[一句话说明已完成什么、当前阻塞什么]'],
        ['责任与时间', '负责人 [真实姓名或未领取]；Agent职责 [实施／复核／未参与]\n复核及签收人 [按需指定]；工作量 [人日]；起止日期 [待填]'],
        ['基线与权威入口', '[唯一项目目录]；[分支／基线版本]；[任务书或接口链接]'],
    ], [1.2, 5.8])
    heading(doc, '一 目标与完成标准', '[本包要解决的问题]。完成后应交付 [可检查对象]，使 [具体消费者] 能够 [下一步动作]。不要只写“完善”“优化”。')
    doc.add_paragraph('范围内 [本包承担的职责与接口]；范围外 [本包不承担的职责]；约束 [时间、设备、数据、软件版本或明确规则]。')
    doc.add_paragraph('启动前确定 AC1 [核心交付的通过标准]、AC2 [下游使用条件]；有明确要求或已发现问题时再加验收项。实际结果填写在第六节，不另建重复验收清单。')
    heading(doc, '二 输入与上下游依赖', '依赖按实际可用性判断；任务状态、实现可用和正式准入分别注明。领取时记录采用的输入版本，上游变更时再判断影响。')
    table(doc, ['上游输入或依赖', '权威位置与版本', '可用性及本包影响'], [
        ['[任务编号／资料]', '[绝对路径或网页链接＋版本]', '[可用／待补；影响什么]'],
        ['[接口／数据／素材]', '[位置＋版本]', '[可用／待补；影响什么]'],
    ], [1.65, 2.85, 2.5])

    doc.add_heading('执行与交付', 0).paragraph_format.page_break_before = True
    heading(doc, '三 过程划分与责任落实', '按可独立完成的步骤正向推进；同时反查每一步是否能汇成目标交付。多人参与时保留一个本包负责人，公共接口或文件明确唯一集成人。')
    table(doc, ['步骤', '行动与产出', '责任人', '完成判断'], [
        ['1 对齐', '[核对输入，确定最小方案]', '[姓名]', '[输入够用，范围明确]'],
        ['2 实现', '[构建最小可检验交付]', '[姓名]', '[产物可打开／可运行]'],
        ['3 验证', '[执行验收与必要复测]', '[姓名]', '[证据对应验收项]'],
        ['4 交接', '[复核、签收、发布与移交]', '[姓名]', '[消费者可直接接续]'],
    ], [.75, 3.15, .85, 2.25])
    heading(doc, '四 实际交付与关键决定', '[用过去时记录实际做过的工作；尚未实施的方案不写成成果。]')
    table(doc, ['交付物', '文件位置或可访问链接', '用途与当前版本'], [
        ['[代码／文档／数据／成品]', '[绝对路径或链接]', '[消费者、用途、版本]'],
        ['[验证产物]', '[位置或链接]', '[本包证明范围]'],
    ], [1.6, 3.2, 2.2])
    doc.add_paragraph('关键决定 [采用的方案]；理由 [与本项目目标的关系]；变更 [相对领取方案改变了什么、谁批准]。没有变更写“无”。')
    heading(doc, '五 必要技能与学习资料', '只列执行本包必须掌握的技能；国内教程或视频用于上手，接口事实优先引用官方文档或原始资料。')
    table(doc, ['技能', '国内教程或视频链接', '学到什么即可开工'], [
        ['[必要技能]', '[标题＋直接链接＋指定章节]', '[与步骤／交付的对应关系]'],
        ['[官方依据，按需]', '[文档或论文链接＋版本]', '[需核实的具体接口或结论]'],
    ], [1.4, 3.3, 2.3])
    doc.add_paragraph('按领域选填，不同时套用全部字段。科研：研究问题、原始资料、主要结果与主张范围。开发：输入输出、运行入口、版本与成品。建模：对象、基线、评价方法与可重跑结果。课程：题目要求、评分点与交付格式。投资研究：数据日期、比较基准与可证伪假设，区分回测和实际结果。')

    doc.add_heading('验收与复核', 0).paragraph_format.page_break_before = True
    heading(doc, '六 验收要求与验证结果', '先写可判定的要求，再填实际结果；证据须能打开或复跑。测试通过不替代真实设备、运行画面、研究结果或外部批准。')
    table(doc, ['验收项', '方法及通过标准', '实际结果与证据'], [
        ['AC1 核心目标', '[动作＋可观察输出或阈值]', '[通过／失败／未执行；位置]'],
        ['AC2 接口交接', '[消费者如何读取或使用]', '[结果；位置]'],
        ['AC3 必要边界', '[任务明确要求或已观测问题]', '[结果；位置]'],
    ], [1.25, 3.05, 2.7])
    doc.add_paragraph('验证环境 [软件版本、输入性质、设备或数据来源]；运行入口 [命令／操作]；结果摘要 [实际数量或指标]；未执行 [具体活动及对签收的影响]。')
    heading(doc, '七 偏离与修复闭环', '只记录已发现的问题、明确未完成项或需求变更；没有问题写“无已发现偏离”。修复后填写复测结果，不以修改完成代替复测通过。')
    table(doc, ['问题与影响', '修复及责任人', '复测与状态'], [
        ['[问题编号；现象；影响哪项验收]', '[修改位置或处置；姓名]', '[结果、证据；关闭／未关闭]'],
    ], [2.45, 2.4, 2.15])
    heading(doc, '八 按需复核与真实签收', '仅在任务要求时填写，否则注明“不适用”。Agent 复核与真人签收分开，不代填姓名、日期、PASS 或 DONE；签收限定到实际审阅的版本与范围。')
    doc.add_paragraph('独立复核人或 Agent [待填]；复核对象 [版本／提交／文件清单]；方式 [文件审阅／复跑／见证]；结论 [待填]；未关闭问题 [待填]。')
    table(doc, ['真实签收字段', '填写内容'], [
        ['签收人及角色', '[真实姓名]；[其在本任务中的责任与权限]'],
        ['对象与依据', '[精确版本或提交]；[本人审阅／本人重跑／见证已有结果]'],
        ['结论及日期', '[待签收／通过／退回／限定通过，按项目规则]；[日期及时区]'],
        ['签收范围', '[本次接受的交付]；[仍未证明的实际运行或外部条件]'],
    ], [1.5, 5.5])

    doc.add_heading('交接与收尾', 0).paragraph_format.page_break_before = True
    heading(doc, '九 下游交接与项目拼合', '逐包验收之后仍要确认总体目标可拼合：输入能对接、输出有人消费、必要活动有人负责。对下游只交付可使用的当前版本。')
    table(doc, ['下游任务与责任人', '交接内容及使用方式', '准入或剩余条件'], [
        ['[编号；姓名或待领取]', '[接口、文件、命令与版本]', '[可开始／需等待什么]'],
        ['[集成任务；责任人]', '[如何组成项目完整交付]', '[整体缺口归哪个任务]'],
    ], [1.8, 3.1, 2.1])
    heading(doc, '十 目录整理与发布记录', '代码与执行材料保留权威入口；Word 面向人工审阅。过时文件按原因归档，引用同步修正；只处理本包范围，不混入其他成员的改动。')
    doc.add_paragraph('人读文档 [位置]；Agent 执行入口 [任务说明、输入、执行、证据、输出的位置]；历史归档 [旧版位置与原因，或无]。')
    doc.add_paragraph('发布记录 [分支]；[提交]；[PR 链接]；[推送状态]；[合并状态]。非代码任务填写实际交付渠道；不适用项写“不适用”。未经授权不提交外部系统或发布全局配置。')
    doc.add_paragraph('联动更新 [项目注册表／README／任务图／进度记录中实际要求的项]；更新结果 [位置与日期]。没有状态变化时，不为形式重复生成。')
    heading(doc, '十一 当前结论与下一包', '[一句话说明：本包交付了什么、证据支持到哪里、还需谁完成什么。]')
    table(doc, ['收尾检查', '实际结论'], [
        ['交付及必要验证', '[完成／未完成；缺什么]'],
        ['独立复核及真实签收', '[各自状态；不得相互替代]'],
        ['发布与下游可用性', '[已交付／已推送／已合并；下游能否使用]'],
        ['下一包', '[推荐编号、领域、任务名；理由与前置条件]'],
    ], [1.8, 5.2])
    doc.add_paragraph('关键来源  [任务书与接口]；[直接支撑结论的原始资料]；[必要教程]。填写可点击的本机文件链接或网页直接链接；不额外复制整份上游材料。')
    doc.core_properties.title = '通用中文逐包任务模板'
    doc.core_properties.subject = '任务定义 执行 验收 签收与交接'
    doc.core_properties.author = ''
    OUT.mkdir(parents=True, exist_ok=True)
    target = OUT / f'{NAME}.docx'
    doc.save(target)
    template = OUT / f'{NAME}.dotx'
    with ZipFile(target) as source, ZipFile(template, 'w', ZIP_DEFLATED) as dest:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == '[Content_Types].xml':
                xml = etree.fromstring(content)
                main = next(x for x in xml if x.get('PartName') == '/word/document.xml')
                main.set('ContentType', 'application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml')
                content = etree.tostring(xml, xml_declaration=True, encoding='UTF-8', standalone=True)
            dest.writestr(item, content)
    print(target)
    print(template)


if __name__ == '__main__':
    build()
