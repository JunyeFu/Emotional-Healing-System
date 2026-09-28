"""Build the reusable Chinese task-package Word form."""
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[3]
OUTPUT = ROOT / 'human/templates/通用中文逐包任务模板_v3.0.docx'


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    widths = (4.0, 13.59) if len(headers) == 2 else (3.9, 7.4, 6.29)
    for column, width in zip(t.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + edge)
        for key, value in {'val': 'single', 'sz': '4', 'color': 'D9D9D9'}.items():
            border.set(qn('w:' + key), value)
        borders.append(border)
    t._tbl.tblPr.append(borders)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'EAF0F4')
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    for row in t.rows:
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        for index, cell in enumerate(row.cells):
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins = OxmlElement('w:tcMar')
            for edge, value in (('top', '85'), ('bottom', '85'), ('left', '110'), ('right', '110')):
                margin = OxmlElement('w:' + edge)
                margin.set(qn('w:w'), value)
                margin.set(qn('w:type'), 'dxa')
                margins.append(margin)
            cell._tc.get_or_add_tcPr().append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(3)
    spacer.add_run().font.size = Pt(3)


def heading(doc, text, prompt):
    doc.add_heading(text, level=1)
    doc.add_paragraph(prompt)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    sec.left_margin = sec.right_margin = Cm(2)
    for name in ('Normal', 'Title', 'Heading 1'):
        s = doc.styles[name]
        s.font.name = 'Microsoft YaHei'
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Heading 1'].font.size = Pt(12)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(9)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(4)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('逐包任务模板 v3.0  |  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.title = '通用中文逐包任务模板'
    doc.core_properties.subject = '任务领取 执行 验收 交接'
    doc.core_properties.author = ''

    doc.add_paragraph('通用中文逐包任务模板', 'Title')
    doc.add_paragraph('一包一份，结论先行。适用于科研、编程、建模、开发和投资研究；随执行增量填写，不把计划写成结果。方括号内容替换为实际信息；无关项写“不适用”，必要时增删表格行。')
    table(doc, ['任务身份', '填写内容'], [
        ('项目与任务', '[项目名称] / [编号] / 【领域】[动作与交付对象]'),
        ('版本与日期', '[文档版本] / [更新日期] / [本次审阅对象版本或提交]'),
        ('责任与领取', '[唯一负责人] / [实施人或Agent] / [复核人或不适用] / [领取日期]'),
        ('入口与状态', '[权威任务文件可点击链接] / [注册状态] / [当前实际进展]'),
    ])
    heading(doc, '一 当前结论', '[用两句话说明本包解决什么、目前完成到哪里，以及本次需要审阅人作出的决定。]')
    doc.add_paragraph('本次请求：[仅知会 / 审阅方案 / 复核结果 / 真实责任人签收]；未关闭事项：[无，或事项与负责人]。')
    doc.add_paragraph('当前状态以项目注册表为准；未有注册表时填写未领取、执行中、待复核或完成，不另建并行状态体系。')
    heading(doc, '二 目标与完成标准', '项目作用：[本包推动哪一项目目标]。范围：[本包承担什么，明确排除什么]。工作量：[预计人日或时间]。')
    table(doc, ['验收编号', '可检查的完成标准', '检查方式'], [
        ('AC-01', '[交付对象、预期行为或结果]', '[运行、复算、文件审阅等]'),
        ('AC-02', '[上游输入与下游接口要求]', '[消费者或交接检查]'),
    ])
    doc.add_paragraph('所需技能与资料：[必需技能、工具及实际版本、国内教学或视频标题和链接]；接口与规范以权威原文为准。资料只列完成本包必需的内容，不填未核验的链接。')
    doc.add_page_break()
    heading(doc, '三 实际交付', '计划时填“拟交付”，完成后改为实物位置；链接应能打开文件或入口，不只写名称。')
    table(doc, ['交付物', '打开或运行入口', '对应验收'], [
        ('[主要交付物]', '[文件链接 / 命令 / 制品入口]', '[AC编号及当前状态]'),
    ])

    heading(doc, '四 执行过程与关键决定', '输入与前置：[上游编号、实际输入位置、版本及是否就绪]。采用方案：[最终选择与一条实质理由]。')
    table(doc, ['步骤与责任人', '动作及阶段出口', '实际进展'], [
        ('1 明确输入 [人]', '[确认需求与最小可检验结果]', '[日期 / 状态 / 结果]'),
        ('2 实施 [人]', '[完成主要交付，对应AC]', '[实际变化或阻断]'),
        ('3 验证与交接 [人]', '[完成必要验证，交给消费者]', '[结果及接收情况]'),
    ])
    doc.add_paragraph('如有方案变化，记录：[原方案 → 当前决定]、[实际理由]及[涉及验收或下游]。未发生变化不补写决策历史。')
    heading(doc, '五 验证与结果', '只记录本次实际执行的检查；未执行写明原因。测试通过、真实运行和研究结论分别表述。')
    table(doc, ['验收项', '实际验证及证据链接', '结果与可支持结论'], [
        ('[AC编号]', '[命令 / 数据 / 截图 / 日期]', '[通过、失败或未验证；结论范围]'),
        ('[AC编号]', '[重跑入口或必要来源]', '[结果、数量或测量值]'),
    ])
    doc.add_paragraph('研究、建模或投资任务按需补充：[数据与分析范围、基线比较、关键假设及结果]。不要用工程完成代替实证，也不要把回测写成未来事实。')
    doc.add_paragraph('最小重跑说明：[工作目录]、[命令及必需输入]、[预期输出]。原始来源、假设、计算结果和推断分别写清，不复制受限数据或凭据。')
    doc.add_page_break()
    heading(doc, '六 偏离与修复闭环', '按任务书、接口和项目目标检查真实偏离；没有发现时写检查范围，不列假想风险。')
    table(doc, ['问题及实际影响', '修复或裁定', '复测与责任'], [
        ('[真实问题或无]', '[修改的权威文件及内容]', '[证据 / 关闭状态 / 负责人]'),
    ])
    doc.add_paragraph('Word中的修订意见应先落实到执行文件，再更新本报告；未关闭项明确责任人与下一动作。')
    heading(doc, '七 目录与文件整理', '当前有效入口：[输入、执行、证据和交付位置]。废弃或历史材料：[保留、归档或删除的对象及理由；无则写无]。')

    heading(doc, '八 上下游交接', '前向检查：下游能否直接消费本包交付。反向检查：本包是否覆盖项目目标要求，是否仍有缺口。')
    table(doc, ['关系', '交接内容与接收人', '当前条件'], [
        ('上游 [编号]', '[实际输入与权威入口]', '[就绪或缺口]'),
        ('下游 [编号]', '[交付、接口及接收人]', '[可开始或仍待何条件]'),
        ('项目总体', '[对目标、研究或产品的具体作用]', '[剩余缺口及承担任务]'),
    ])
    doc.add_paragraph('交接判据：[消费者可以直接打开、运行或继续工作的具体条件]。未满足时指定承担后续动作的任务与责任人，不凭上游状态自动解锁。')
    doc.add_page_break()
    heading(doc, '九 审阅与责任签收', '独立Agent复核与真实责任人签收分开记录；仅在任务要求时开展。未签收填“待签收”；不得代填姓名、日期或结论。')
    table(doc, ['审阅环节', '实际记录'], [
        ('独立复核', '[复核人或Agent / 日期 / 对象版本 / 结论 / 报告链接]'),
        ('真实责任人签收', '[姓名 / 日期 / PASS或退回 / 签收范围 / 原始记录]'),
        ('真实活动或外部条件', '[设备、现场、许可或平台回执；仅任务要求时填写]'),
    ])
    doc.add_paragraph('完成层次：[实现] / [必要验证] / [独立复核] / [责任签收] / [发布或研究准入]。逐项填写真实状态，后项不由前项自动推导。')
    heading(doc, '十 收尾与下一包', '收尾：[交付位置；任务要求时填写commit、push、PR、合并或锁定版本]。')
    doc.add_paragraph('下一包：[编号、领域与名称]；领取人：[姓名或未领取]；启动条件：[已满足或具体缺口]。下一必要动作：[一个明确动作及负责人]。')
    heading(doc, '十一 审阅来源', '任务权威：[链接]；执行与重跑入口：[链接]；本次证据：[链接]；历史签署：[链接或不适用]。来源只列支撑本包判断的材料，当前结果与历史结果分开。')
    doc.add_paragraph('使用约定：小任务保留目标、交付、验证与责任即可压缩为一页；复杂任务增加必要附件，不逐日重复叙述。同一事实只写一次，其他位置引用。')
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    template = OUTPUT.with_suffix('.dotx')
    content_types = '{http://schemas.openxmlformats.org/package/2006/content-types}'
    with ZipFile(OUTPUT) as source, ZipFile(template, 'w', ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == '[Content_Types].xml':
                tree = ElementTree.fromstring(data)
                for entry in tree.findall(content_types + 'Override'):
                    if entry.get('PartName') == '/word/document.xml':
                        entry.set('ContentType', 'application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml')
                data = ElementTree.tostring(tree, encoding='utf-8', xml_declaration=True)
            target.writestr(item, data)
    print(OUTPUT)
    print(template)


if __name__ == '__main__':
    main()
