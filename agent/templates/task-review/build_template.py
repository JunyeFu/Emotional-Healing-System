"""Build a portable Chinese task-review Word template."""
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


def add_table(doc, headings, rows, widths):
    table = doc.add_table(rows=1, cols=len(headings))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for cell, heading, width in zip(table.rows[0].cells, headings, widths):
        cell.width = Inches(width)
        cell.text = heading
    for row in rows:
        cells = table.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = Inches(width)
            cell.text = value
    for index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            properties = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'bottom', 'left', 'right'):
                border = OxmlElement('w:' + edge)
                for name, value in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
                    border.set(qn('w:' + name), value)
                borders.append(border)
            properties.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'bottom', 'left', 'right'):
                margin = OxmlElement('w:' + edge)
                margin.set(qn('w:w'), '90')
                margin.set(qn('w:type'), 'dxa')
                margins.append(margin)
            properties.append(margins)
            if index == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'E8EEF4')
                properties.append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.1
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    run.bold = index == 0
    return table


def section(doc, title, texts):
    doc.add_heading(title, level=1)
    for text in texts:
        doc.add_paragraph(text)


def build(output):
    doc = Document()
    page = doc.sections[0]
    page.page_width, page.page_height = Inches(8.5), Inches(11)
    page.top_margin = page.bottom_margin = Inches(0.65)
    page.left_margin = page.right_margin = Inches(0.75)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.15
    doc.styles['Title'].font.size = Pt(22)
    doc.styles['Heading 1'].font.size = Pt(13)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)
    header = page.header.paragraphs[0]
    header.text = '逐包任务总结    人类审阅版本'
    header.style = doc.styles['Normal']
    footer = page.footer.paragraphs[0]
    footer.text = '通用中文模板 v1.0    '
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    doc.add_paragraph('逐包任务总结', 'Title')
    doc.add_paragraph('项目［项目名称］    任务［编号与领域］    当前文档［版本与日期］')
    section(doc, '一 当前结论', [
        '［用两至三句说明解决了什么、当前是否达到验收、还有哪项需要审阅或完成。］',
        '本次审阅动作：［签收／修复后复审／只阅知／需作出选择；写明理由。］'
    ])
    add_table(doc, ['项目', '实际登记'], [
        ('任务名称与领域', '［编号、名称、领域，名称与任务注册表一致］'),
        ('责任人及复核人', '［领取人／实施Agent／独立复核人／真实第二人］'),
        ('任务与交付状态', '［当前注册状态；实现、验证、签收分别说明］'),
        ('当前权威位置', '［Agent执行目录、代码或设计主文件；使用可打开位置］')
    ], [1.65, 5.35])
    section(doc, '二 目标与完成标准', [
        '任务要解决的问题：［说明本包在项目主线中解决的具体问题。］',
        '验收要求：［引用任务书AC；写出怎样的结果算完成。］',
        '本包范围与变更：［实际承担内容；仅写影响验收的范围调整及裁定来源。］'
    ])
    section(doc, '三 实际交付', [
        '［说明交付物做什么、在哪里、如何打开或运行，给出最有代表性的结果。］'
    ])
    add_table(doc, ['交付物', '实际位置', '验收对应'], [
        ('［交付名称］', '［文件或入口］', '［AC编号及结果］'),
        ('［交付名称］', '［文件或入口］', '［AC编号及结果］')
    ], [1.8, 3.8, 1.4])

    doc.add_page_break()
    section(doc, '四 执行过程与关键决定', [
        '输入与前提：［列出本次真正使用的输入、设备、数据或上游结果。］',
        '推进过程：［按实际步骤写输入读取、实现、检查和收尾；不要复制未执行计划。］',
        '关键决定：［写最终采用的方案、直接理由与项目约束；有必要才记录替代方案。］'
    ])
    section(doc, '五 验证与结果', [
        '［先说明验证结果和它支持的完成范围，未执行项写未执行及原因。］'
    ])
    add_table(doc, ['验收项', '检查与证据', '结果及范围'], [
        ('［AC或检查］', '［命令、截图、数据或真实回执位置］', '［通过／失败／待验证］'),
        ('［AC或检查］', '［检查对象、时间与结果位置］', '［能够支持的结论］')
    ], [1.25, 3.7, 2.05])
    section(doc, '六 偏离与修复闭环', [
        '［检查是否偏离任务验收、当前设计、接口或项目主线；没有发现时说明检查范围。］'
    ])
    add_table(doc, ['问题与影响', 'Agent层修复', '复测与状态'], [
        ('［事实问题；影响哪个下游］', '［修复文件与实际改变］', '［复测结果；关闭或未关闭］'),
        ('［无则写无已发现项］', '［不适用或处理位置］', '［真实状态］')
    ], [2.4, 2.5, 2.1])
    doc.add_paragraph('Word回填：［确认修复已发生在权威执行文件，再写入当前总结；列明未关闭项的负责人和下一步。］')
    section(doc, '七 目录与文件整理', [
        '当前文件：［执行、证据、输出各自位置；确认没有未归类文件。］',
        '过期文件处理：［删除或归档的文件、替代位置及保留理由；没有则写无。］'
    ])

    doc.add_page_break()
    section(doc, '八 上下游串联', [
        '项目主线位置：［此包如何推进整体目标，当前交付是否改变原研究或产品方向。］',
        '上游实际输入：［上游编号、真实交付、版本要求及缺口。］',
        '下游交接：［下游编号与责任人，需要接收什么、用在哪一步、怎样验证接入。］'
    ])
    add_table(doc, ['关系', '交接内容', '当前状态'], [
        ('［上游或下游编号］', '［实际输入或输出及权威位置］', '［已就绪／缺口及责任人］'),
        ('［跨包影响］', '［变更影响与需要更新的消费者］', '［处理状态］')
    ], [1.7, 3.8, 1.5])
    section(doc, '九 审阅与责任签收', [
        '独立复核：［真实复核人或Agent、对象、日期、结论及问题位置。］',
        '真实第二人签收：［姓名、日期、结论与签收范围；未签收时填写待签收。］',
        '需实际活动证明的事项：［有此要求时填写设备、现场、外部批准或参与者活动证据；没有则写不适用。］',
        '当前总状态：［实现／验证／复核／签收／活动／发布按实际分别说明。］'
    ])
    section(doc, '十 收尾与下一包', [
        '发布记录：［本任务要求的commit、push、PR、合并或交付位置；未要求则写不适用。］',
        '下一包：［编号、名称、可领取条件、负责人；写清是否已经解锁。］',
        '继续动作：［最先要执行的一步与所需输入；有阻断时说明具体阻断。］'
    ])
    section(doc, '十一 审阅来源', [
        '结构化总结：［Agent层当前summary文件位置］',
        '证据与签署：［本轮结果及历史签署位置；历史与本轮证据分别列出］',
        '串联文档：［项目任务串联审阅文件位置］'
    ])
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == '__main__':
    root = Path(__file__).resolve().parents[3]
    build(root / 'human/templates/通用中文逐包任务总结模板.docx')
