"""Generate current human reviews from task execution summaries."""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
GOV = ROOT / '00-项目管理/01-项目章程与规划/2026-08-05_SRP_IJHCI_全项目1-12步规划设计包_v1.0/24_团队任务与项目治理'


def read_json(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))


def document(title):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.styles['Normal'].paragraph_format.space_after = Pt(7)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.15
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.styles['Title'].font.size = Pt(21)
    doc.styles['Heading 1'].font.size = Pt(13)
    doc.add_paragraph(title, 'Title')
    footer = section.footer.paragraphs[0]
    footer.add_run('SRP 任务包审阅  ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    return doc


def render_summary(summary, output):
    doc = document(summary['title'])
    doc.add_paragraph(f"审阅日期 {summary['review_date']}    任务 {summary['task_id']}")
    doc.add_paragraph(summary['conclusion'])
    acceptance = summary['historical_acceptance']
    doc.add_heading('历史签收对象', level=1)
    doc.add_paragraph(f"审核人 {acceptance['reviewer']}    日期 {acceptance['date']}")
    doc.add_paragraph(f"审核提交 {acceptance['commit']}")
    doc.add_paragraph(acceptance['method'])
    for section in summary['sections']:
        doc.add_heading(section['heading'], level=1)
        for paragraph in section['paragraphs']:
            doc.add_paragraph(paragraph)
    doc.add_heading('审阅来源', level=1)
    doc.add_paragraph(f"Agent层总结 agent/tasks/{summary['task_id']}/outputs/summary.json")
    doc.add_paragraph(f"验证记录 agent/tasks/{summary['task_id']}/evidence/verification.json")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def check_task(task_id, registry):
    if task_id not in registry:
        raise ValueError(f'Unknown task: {task_id}')
    package = ROOT / 'agent/tasks' / task_id
    required = ('TASK.md', 'inputs/sources.json', 'execution/verify.ps1',
                'evidence/verification.json', 'outputs/summary.json', 'archive/README.md')
    for path in required:
        if not (package / path).is_file():
            raise ValueError(f'Missing package file: {task_id}/{path}')
    allowed_dirs = {'inputs', 'execution', 'evidence', 'outputs', 'archive'}
    for entry in package.iterdir():
        if entry.name != 'TASK.md' and entry.name not in allowed_dirs:
            raise ValueError(f'Unclassified package file: {entry}')
    for source in read_json(package / 'inputs/sources.json')['paths']:
        resolved = (ROOT / source).resolve()
        if not resolved.is_relative_to(ROOT) or not resolved.exists():
            raise ValueError(f'Missing source: {source}')
    summary = read_json(package / 'outputs/summary.json')
    evidence = read_json(package / 'evidence/verification.json')
    if summary['task_id'] != task_id or evidence['task_id'] != task_id:
        raise ValueError('Task identity mismatch')
    if not evidence['checks'] or any(c['exit_code'] != 0 for c in evidence['checks']):
        raise ValueError(f'Unresolved verification failure: {task_id}')
    return summary


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('task_id')
    parser.add_argument('--check', action='store_true')
    args = parser.parse_args()
    with (GOV / '05_可领取任务包.csv').open(encoding='utf-8-sig', newline='') as stream:
        rows = list(csv.DictReader(stream))
    registry = {r['task_id']: r for r in rows}
    summary = check_task(args.task_id, registry)
    if args.check:
        review = ROOT / 'human/tasks' / args.task_id / 'summary.docx'
        if review.is_file():
            text = '\n'.join(p.text for p in Document(review).paragraphs)
            expected = [summary['conclusion'], summary['historical_acceptance']['commit']]
            expected.extend(p for s in summary['sections'] for p in s['paragraphs'])
            if any(value not in text for value in expected):
                raise ValueError(f'Word content drift: {args.task_id}')
        print(f'PASS package structure and sources: {args.task_id}')
        return
    render_summary(summary, ROOT / 'human/tasks' / args.task_id / 'summary.docx')
    chain = document('SRP 任务包串联审阅')
    chain.add_paragraph('本文件由已整理包的Agent层总结串联生成。用于检查任务完成范围和下游交接；未整理包不推定已核验。')
    completed = []
    findings = []
    for row in rows:
        path = ROOT / 'agent/tasks' / row['task_id'] / 'outputs/summary.json'
        if not path.exists():
            continue
        item = check_task(row['task_id'], registry)
        completed.append(row['task_id'])
        chain.add_heading(f"{row['task_id']} {row['domain']}", level=1)
        chain.add_paragraph(item['conclusion'])
        handoff = next(s for s in item['sections'] if s['heading'] == '上下游交接')
        for paragraph in handoff['paragraphs']:
            chain.add_paragraph(paragraph)
        for finding in item['findings']:
            findings.append({'task_id': row['task_id'], **finding})
            chain.add_paragraph(f"{finding['id']} {finding['status']}：{finding['description']}")
    chain.add_heading('待整理队列', level=1)
    remaining = [r['task_id'] for r in rows if r['task_id'] not in completed]
    chain.add_paragraph(f"已整理{len(completed)}包，待整理{len(remaining)}包。")
    chain.add_paragraph('、'.join(remaining))
    chain.add_paragraph('根目录仍有旧业务目录，最终双层迁移未完成。下一包按上述队列顺序推进。')
    chain.save(ROOT / 'human/project-review.docx')
    progress = {'completed_packages': completed, 'remaining_packages': remaining,
                'root_migration_complete': False, 'findings': findings}
    (ROOT / 'agent/normalization-progress.json').write_text(
        json.dumps(progress, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'WROTE {args.task_id} Word and chained review; {len(remaining)} packages remain')


if __name__ == '__main__':
    main()
